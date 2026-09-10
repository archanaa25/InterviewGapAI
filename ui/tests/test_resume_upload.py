"""Tests for format-aware resume extraction owned by the UI boundary."""

from __future__ import annotations

import io

import pytest
from docx import Document

from ui.resume_upload import (
    MAX_RESUME_BYTES,
    ResumeFormat,
    ResumeUploadError,
    extract_resume,
)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("John Mathew", level=1)
    document.add_paragraph("Python engineer with applied RAG experience.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    )
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{number} 0 obj\n".encode("ascii"))
        content.extend(value)
        content.extend(b"\nendobj\n")
    xref = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(content)


@pytest.mark.parametrize(
    ("filename", "media_type", "data", "expected", "phrase"),
    [
        ("resume.md", "text/markdown", b"# John\nPython", ResumeFormat.MARKDOWN, "Python"),
        (
            "resume.csv",
            "text/csv",
            b"section,detail\nskills,Python and RAG\n",
            ResumeFormat.CSV,
            "skills | Python and RAG",
        ),
        (
            "resume.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx_bytes(),
            ResumeFormat.DOCX,
            "John Mathew",
        ),
        (
            "resume.pdf",
            "application/pdf",
            _pdf_bytes("John uses Python and RAG"),
            ResumeFormat.PDF,
            "Python and RAG",
        ),
    ],
)
def test_supported_modern_formats_extract_text(
    filename: str,
    media_type: str,
    data: bytes,
    expected: ResumeFormat,
    phrase: str,
) -> None:
    result = extract_resume(filename, data, media_type)

    assert result.format is expected
    assert phrase in result.text
    assert result.candidate_id.startswith("upload-")


def test_legacy_doc_uses_injected_converter() -> None:
    content = bytes.fromhex("d0cf11e0a1b11ae1") + b"synthetic"

    result = extract_resume(
        "resume.doc",
        content,
        "application/msword",
        legacy_doc_converter=lambda _data: "John Mathew\nPython engineer",
    )

    assert result.format is ResumeFormat.DOC
    assert result.warnings


@pytest.mark.parametrize(
    ("filename", "data", "media_type", "code"),
    [
        ("resume.txt", b"hello", "text/plain", "unsupported_format"),
        ("resume.pdf", b"not pdf", "application/pdf", "invalid_signature"),
        ("resume.csv", b"name,value", "application/pdf", "media_type_mismatch"),
        ("resume.md", b"\xff", "text/markdown", "invalid_encoding"),
        ("resume.md", b"x" * (MAX_RESUME_BYTES + 1), "text/markdown", "file_too_large"),
    ],
)
def test_invalid_uploads_fail_closed(
    filename: str,
    data: bytes,
    media_type: str,
    code: str,
) -> None:
    with pytest.raises(ResumeUploadError) as captured:
        extract_resume(filename, data, media_type)

    assert captured.value.code == code
